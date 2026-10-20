"""The Excel files written by DuckDB must contain exactly the values of the parquet read tables."""
import shutil

import duckdb
import openpyxl
import pandas as pd
import pytest

from apscale import j_generate_read_table as j


def _read_parts(folder, prefix):
    frames = []
    for part in sorted(folder.glob(f"{prefix}_read_table_part_*.xlsx"), key=lambda f: int(f.stem.split("_")[-1])):
        wb = openpyxl.load_workbook(part, read_only=True)
        assert wb.sheetnames == ["Sheet1"]
        frames.append(pd.DataFrame(list(wb.active.values)[1:], columns=list(next(wb.active.iter_rows(max_row=1, values_only=True)))))
        wb.close()
    return frames


@pytest.mark.parametrize("sequence_type, prefix", [("sequence", "0_t_sequence"), ("group", "1_t_group")])
def test_excel_parts_equal_parquet(storage, tmp_path, monkeypatch, sequence_type, prefix):
    db = tmp_path / "db.duckdb"
    shutil.copy(storage, db)
    out, temp = tmp_path / "out", tmp_path / "temp"
    out.mkdir(), temp.mkdir()
    monkeypatch.setattr(j, "EXCEL_MAX_CELLS", 400)  # few samples per part -> many parts
    j.generate_read_table("t", out, db, sequence_type, temp)

    parts = _read_parts(out, prefix)
    assert len(parts) > 3
    parquet = duckdb.sql(
        f"SELECT * FROM read_parquet({j.quote_literal(next(out.glob(f'{prefix}_read_table.parquet.snappy')))})"
    ).df()

    seen = []
    for part in parts:
        # same sequences in the same order in every part, identical to the parquet
        assert part["hash"].tolist() == parquet["hash"].tolist()
        assert part["sequence"].tolist() == parquet["sequence"].tolist()
        samples = [c for c in part.columns if c not in ("hash", "sequence")]
        seen += samples
        assert part[samples].astype("int64").equals(parquet[samples].astype("int64"))
    assert sorted(seen) == sorted(c for c in parquet.columns if c not in ("hash", "sequence"))
